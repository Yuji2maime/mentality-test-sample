import streamlit as st
import google.generativeai as genai
import datetime
import json
import os
import pandas as pd
import plotly.graph_objects as go
st.set_page_config(page_title="AI統合型 認知特性テスト", layout="wide")

# セッション状態の初期化
if "submissions" not in st.session_state:
    st.session_state.submissions = []
if "view_mode" not in st.session_state:
    st.session_state.view_mode = "applicant"

# SecretsからAPIキーを自動読み込み
api_key = st.secrets["GEMINI_API_KEY"]
# --- サイドバー機能：ホーム画面追加案内 & データ削除 ---
with st.sidebar:
    # 免責・注意事項の表示（折りたたみ式）
    with st.sidebar.expander("⚠️ ご利用にあたっての重要事項（免責・禁止事項）", expanded=False):
        st.warning("""
**【1. ツールの目的と運用の限界】**
本ツールは、応募者や従業員の思考・コミュニケーション特性を把握し、配置のミスマッチや組織内摩擦の低減、日常の労務管理（ハラスメントリスクの低減等）を補助するためのものです。**医学的な診断を目的とするものではなく、面接時におけるメンタル疾患の発見・特定等を目的とした使用はできません。**

**【2. 免責事項（不可抗力と危険負担）】**
予期せぬシステムエラーやバグ、通信障害、第三者による不正アクセス（ハッキング等）、および天災地変等の不可抗力（自然災害等）により本システムが停止・誤作動・情報漏洩した場合、それに伴う利用者のいかなる損害についても開発者は責任を負いかねます。
**【3. 転売・譲渡の禁止と法的措置】**
開発者の事前の同意なく、本ツールのプログラム、URL、出力結果のフォーマット等を複製、無断転売、譲渡、貸与することを固く禁じます。違反行為が発覚した場合、損害賠償等の民事上の措置に加え、直ちに**刑事告訴**等の厳格な法的措置を講じます。
        """)
    st.header("⚙️ アプリ機能")

    # 1. ホーム画面・デスクトップ追加案内
    with st.expander("📱 ホーム画面・デスクトップに追加"):
        st.markdown("""
        次回からワンタップで開けるよう、追加しておくと便利です。
        
        **【iPhone (Safari)】**
        1. 画面下の **共有ボタン（□に↑）** をタップ
        2. **「ホーム画面に追加」** を選択
        
        **【Android (Chrome)】**
        1. 画面右上の **「3点リーダー（⋮）」** をタップ
        2. **「ホーム画面に追加」** または **「アプリをインストール」** を選択
        
        **【パソコン (Chrome / Edge)】**
        1. 画面右上の **「3点リーダー（︙ または …）」** をクリック
        2. Chrome: **「保存して共有」** ＞ **「ショートカットを作成」**
        3. Edge: **「アプリ」** ＞ **「このサイトをアプリとしてインストール」**
        """)

    # 2. 保存データの削除機能
    st.subheader("🗑️ データ管理")
    if st.button("保存された履歴データをすべて削除", use_container_width=True):
        st.session_state.submissions = []
        st.success("すべての履歴データを削除しました！")
        st.rerun()
# --- データ管理エリアへのCSVダウンロード機能追加 ---
st.sidebar.markdown("---")
st.sidebar.subheader("📥 データダウンロード")

if "submissions" in st.session_state and st.session_state.submissions:
    # データをPandasのDataFrameに変換
    df = pd.DataFrame(st.session_state.submissions)

    # すべての列のカッコ [ ] や引用符 ' ' を綺麗に外す処理
    for col in df.columns:
        df[col] = df[col].apply(lambda x: ', '.join(x) if isinstance(x, list) else x)
        df[col] = df[col].apply(lambda x: str(x).replace("['", "").replace("']", "") if isinstance(x, str) and str(x).startswith("['") else x)

    # 綺麗にしたデータをCSVにエクスポート
    csv_data = df.to_csv(index=False).encode("utf-8-sig")
    st.sidebar.download_button(
        label="📄 履歴をCSVでダウンロード",
        data=csv_data,
        file_name="cognitive_test_submissions.csv",
        mime="text/csv",
    )# === ▼ ここから追加：検索メニューと表の表示 ▼ ===
    st.sidebar.markdown("---")
    st.sidebar.subheader("🔍 データの検索・絞り込み")
    
    # サイドバー：キーワード検索
    search_query = st.sidebar.text_input("キーワード検索 (名前やメモなど)")
    
    # サイドバー：主タイプの絞り込み
    type_options = ["Fe-Si", "Se-Ti", "Ne-Fi", "Ni-Te", "Si-Fe", "Ti-Ne", "Fi-Ne", "Te-Ni", "その他"]
    selected_types = st.sidebar.multiselect("主タイプで絞り込み", type_options)
    
    # 絞り込みの実行
    filtered_df = df.copy()
    if search_query:
        mask = filtered_df.astype(str).apply(lambda x: x.str.contains(search_query, case=False, na=False)).any(axis=1)
        filtered_df = filtered_df[mask]
    if selected_types:
        if "主タイプ" in filtered_df.columns:
            filtered_df = filtered_df[filtered_df["主タイプ"].isin(selected_types)]
            
    # メイン画面：絞り込まれたデータを表として表示
    st.write("### 📋 提出データ一覧")
    st.dataframe(filtered_df, use_container_width=True)
    # === ▲ 追加ここまで ▲ ===
else:
    st.sidebar.info("ダウンロード可能なデータはありません。 ")
st.sidebar.markdown("---")
# 選択肢の定義
MAIN_TYPE_OPTIONS = ["Fe-Si", "Se-Ti", "Ne-Fi", "Ni-Te", "Si-Fe", "Ti-Ne", "Fi-Ne", "Te-Ni", "その他"]
AUX_FUNC_OPTIONS = ["外向感情(Fe)", "内向感覚(Si)", "外向直観(Ne)", "内向思考(Ti)", "外向感覚(Se)", "内向感情(Fi)", "外向思考(Te)", "内向直観(Ni)"]

def analyze_text_with_ai(text, key):
    """応募者の文章をGemini APIで自動解析する関数（数値抽出を強化）"""
    if not key:
        return [], [], "※APIキー未設定のため自動解析ステップ。手動で入力してください。", {}

    try:
        genai.configure(api_key=key)

        prompt = f"""
以下の応募者の記述文章をプロの労務・人事評価者の視点から客観的かつ厳格に分析してください。

【応募者記述文章】
{text}

【選択肢の定義】
MAIN_TYPE_OPTIONS = {MAIN_TYPE_OPTIONS}
AUX_FUNC_OPTIONS = {AUX_FUNC_OPTIONS}

【分析指示】
1. 主タイプ（該当するもの）：MAIN_TYPE_OPTIONSの中から1つ以上選んでください。
2. 補助機能（複数認定）：AUX_FUNC_OPTIONSの中から選んでください。
3. 採用・評価メモ：以下の観点を含め、客観的・事実ベースのトーン（150〜250文字程度）でまとめてください。
   - 組織適応性および規律・コンプライアンス意識
   - 思考・行動の偏りと潜在的リスク（例：独断的傾向、対人コミュニケーションの課題、ストレス耐性など）
   - 評価時の注意事項・面接での確認推奨事項
4. 認知特性スコア（0〜100の数値）：記述内容に基づき以下の5項目を評価・数値化してください。
   - logic: 論理的分析力
   - intuition: 直観・本質把握
   - planning: 計画・規律性
   - independence: 独立・内省力
   - flexibility: 対人・柔軟性

必ず以下のJSON形式のみで出力してください。Markdownの装飾コードブロック（```json ... ```）は含めないでください。

{{
  "main_types": ["タイプ名"],
  "aux_funcs": ["補助機能1", "補助機能2"],
  "eval_memo": "評価メモ本文...",
  "scores": {{
    "logic": 70,
    "intuition": 60,
    "planning": 80,
    "independence": 75,
    "flexibility": 50
  }}
}}
"""

        available_models = []
        try:
            for m in genai.list_models():
                if 'generateContent' in getattr(m, 'supported_generation_methods', []):
                    available_models.append(m.name)
        except Exception as e:
            return [], [], f"APIキー認証エラー: {e}", {}

        if not available_models:
            available_models = ["models/gemini-1.5-flash", "models/gemini-2.0-flash", "models/gemini-1.5-pro"]

        priority_keywords = ["2.0-flash", "1.5-flash", "flash", "1.5-pro"]
        sorted_models = []
        for kw in priority_keywords:
            for m in available_models:
                if kw in m and m not in sorted_models:
                    sorted_models.append(m)
        for m in available_models:
            if m not in sorted_models:
                sorted_models.append(m)

        res_text = None
        last_error = None

        for model_name in sorted_models:
            try:
                try:
                    model = genai.GenerativeModel(
                        model_name,
                        generation_config={"response_mime_type": "application/json"}
                    )
                    response = model.generate_content(prompt)
                except Exception:
                    model = genai.GenerativeModel(model_name)
                    response = model.generate_content(prompt)

                if response and response.text:
                    res_text = response.text.strip()
                    break
            except Exception as err:
                last_error = err
                continue

        if not res_text:
            raise last_error if last_error else Exception("利用可能なGeminiモデルで応答が取得できませんでした。")

        if "```json" in res_text:
            res_text = res_text.split("```json")[1].split("```")[0].strip()
        elif "```" in res_text:
            res_text = res_text.split("```")[1].split("```")[0].strip()

        data = json.loads(res_text)
        
        # スコアの型変換と安全性の確保
        raw_scores = data.get("scores", {})
        scores = {
            "logic": int(raw_scores.get("logic", 50)),
            "intuition": int(raw_scores.get("intuition", 50)),
            "planning": int(raw_scores.get("planning", 50)),
            "independence": int(raw_scores.get("independence", 50)),
            "flexibility": int(raw_scores.get("flexibility", 50))
        }

        return data.get("main_types", []), data.get("aux_funcs", []), data.get("eval_memo", ""), scores

    except Exception as e:
        return [], [], f"AI解析エラー: {e}", {}
# 画面切り替えボタン
col_nav1, col_nav2 = st.columns([1, 1])
with col_nav1:
    if st.button("応募者用画面を表示"):
        st.session_state.view_mode = "applicant"
        st.rerun()
with col_nav2:
        with st.expander("採用側（管理）画面へ"):
            admin_pass = st.text_input("パスワードを入力", type="password")
            if admin_pass == "7777":  # 好きなパスワードに変更可能です
                if st.button("ログインして切り替え"):
                    st.session_state.view_mode = "admin"
                    st.rerun()
            elif admin_pass != "":
                st.error("パスワードが違います")

st.divider()

# 1. 応募者画面
if st.session_state.view_mode == "applicant":
    st.title("思考・表現力セッション")
    st.write("政治・経済、趣味、恋愛など、あなたが今最も関心のあることや語りたいテーマについて、制限時間内に自由に記述してください。納得した時点でいつでも終了できます。")
    
    user_input = st.text_area("記述欄", height=200, key="applicant_text")
    
if st.button("これで完了する（終了）", type="primary"):
        if user_input.strip():
            with st.spinner("AIが回答内容を事前解析中..."):
                mains, auxs, memo, scores = analyze_text_with_ai(user_input, api_key)

            new_data = {
                "id": len(st.session_state.submissions) + 1,
                "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "text": user_input,
                "selected_mains": mains,
                "selected_auxs": auxs,
                "memo": memo,
                "scores": scores
            }
            st.session_state.submissions.append(new_data)
            st.success("送信が完了しました。ご協力ありがとうございました。")
        else:
            st.warning("文章を入力してから送信してください。")
# 2. 管理者画面
else:
    st.title("採用管理画面")
    
    # ★ 免責・運用原則テキスト
    st.warning("※本解析結果は思考傾向の示唆に留まる補助情報であり、適正な採用・配属を保証するものではありません。最終的な採用・配属決定は面接や総合評価に基づき、担当者ご自身の責任で行ってください。")
    
    if not st.session_state.submissions:
        st.info("まだ提出されたデータはありません。")
    else:
       for idx, sub in enumerate(st.session_state.submissions):
        with st.expander(f"提出データ #{sub['id']} (日時: {sub['timestamp']})", expanded=True):
            # --- 追加：個別削除ボタン ---
            if st.button("🗑️ このデータを削除", key=f"del_btn_{sub['id']}"):
                st.session_state.submissions.pop(idx)
                st.rerun()
            # ----------------------------
            
            st.subheader("【応募者の記述内容】")
            st.write(sub["text"])
            
            st.subheader("【AI拡張解析・ラベリングエリア】")
            col1, col2 = st.columns(2)
            
            with col1:
                selected_mains = st.multiselect(
                    "主タイプ (該当するものすべて) ",
                    options=MAIN_TYPE_OPTIONS,
                    default=sub["selected_mains"],
                    key=f"main_{sub['id']}"
                )
            with col2:
                selected_auxs = st.multiselect(
                    "補助機能 (複数認定) ",
                    options=AUX_FUNC_OPTIONS,
                    default=sub["selected_auxs"],
                    key=f"aux_{sub['id']}"
                )
                
                memo = st.text_area(
                    "採用・評価メモ（認知の癖、リスク、矛盾点など）",
                    value=sub["memo"],
                    height=120,
                    key=f"memo_{sub['id']}"
                )
                
            if st.button("💾 評価を保存する", key=f'save_{sub["id"]}', type="primary"):
                    # 画面中央に風船を飛ばす！🎈
                    st.balloons()
                    
                    # 保存データを更新
                    st.session_state.submissions[idx]["selected_mains"] = selected_mains
                    st.session_state.submissions[idx]["selected_auxs"] = selected_auxs
                    st.session_state.submissions[idx]["memo"] = memo
                    
                    # 成功メッセージを表示
                    st.success(f"🎉 提出データ #{sub['id']} の評価を保存・更新しました！") 
# --- 診断結果のレーダーチャート表示（UIテスト用） ---
st.subheader("📊 認知特性・傾向分析（レーダーチャート）")

# 評価軸（項目）の設定
categories = ['論理的分析力', '直観・本質把握', '計画・規律性', '独立・内省力', '対人・柔軟性']

# 診断結果から各評価軸のスコアを取得
# エラー回避：保存データがあれば最新のものを取得、なければ空にする
current_sub = st.session_state.submissions[-1] if st.session_state.submissions else {}
ai_scores = current_sub.get('scores', {})
scores = [
    ai_scores.get('logic', 50),
    ai_scores.get('intuition', 50),
    ai_scores.get('planning', 50),
    ai_scores.get('independence', 50),
    ai_scores.get('flexibility', 50)
]

# レーダーチャートの生成
fig = go.Figure()
fig.add_trace(go.Scatterpolar(
r=scores + [scores[0]], 
theta=categories + [categories[0]],
fill='toself',
fillcolor='rgba(0, 123, 255, 0.3)',
line=dict(color='rgba(0, 123, 255, 1.0)', width=2),
name='特性スコア'
))

fig.update_layout(
polar=dict(
radialaxis=dict(visible=True, range=[0, 100])
),
showlegend=False,
margin=dict(l=20, r=20, t=20, b=20)
)

st.plotly_chart(fig, use_container_width=True)
# === ここから下を新しく追加 ===
st.markdown("---")
st.subheader("💾 データのダウンロード")

# === Excelファイルの作成とレイアウト設定 ===
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = "評価結果"

# ヘッダー（見出し）の設定
headers = ["応募者ID", "主タイプ", "補助機能", "採用・評価メモ"]
ws.append(headers)

# ヘッダーの装飾（青背景、白文字、太字、中央揃え）
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
for col_num, cell in enumerate(ws[1], 1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# データの追加
for sub in st.session_state.submissions:
    sub_id = str(sub.get('id', ''))
    mains = "、".join(sub.get('selected_mains', []))
    auxs = "、".join(sub.get('selected_auxs', []))
    memo = str(sub.get('memo', ''))
    ws.append([sub_id, mains, auxs, memo])

# 列幅の明示的設定（日本語が途切れないよう広めに確保）
ws.column_dimensions['A'].width = 15  # 応募者ID
ws.column_dimensions['B'].width = 15  # 主タイプ
ws.column_dimensions['C'].width = 40  # 補助機能
ws.column_dimensions['D'].width = 60  # 採用・評価メモ

# データ行の配置設定（上揃え・メモ欄の折り返し）
for row in ws.iter_rows(min_row=2):
    for cell in row:
        if cell.column_letter == 'D':
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            cell.alignment = Alignment(vertical="top")
# Excelファイルをメモリ上に保存
excel_buffer = io.BytesIO()
wb.save(excel_buffer)
excel_data = excel_buffer.getvalue()

st.download_button(
    label="📥 評価結果をExcelでダウンロード",
    data=excel_data,
    file_name="evaluation_results.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
# ========================================

# === ここまで ===
st.markdown("---")
if st.button("🔄 次の人のテストを始める（画面リセット）"):
    # 保存済みのCSVデータ(submissions)以外をクリアする
    for key in list(st.session_state.keys()):
        if key != "submissions":
            del st.session_state[key]
    st.rerun()
